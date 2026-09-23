from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


OUT_DIR = Path(__file__).resolve().parent
OUT_XLSX = OUT_DIR / "TRIBEv2_literature_audit_2026-07-21.xlsx"


search_plan = [
    {
        "Step": 1,
        "Search Area": "Exact TRIBE v2 paper trail",
        "Queries / Methods": '"TRIBE v2" fMRI; "TRIBEv2" fMRI; "A foundation model of vision, audition, and language"; "facebook/tribev2"; arXiv and Meta pages',
        "Purpose": "Find the foundation paper, code, official claims, model limits, and version history.",
        "Inclusion Rule": "Include primary TRIBE v2 paper, official code/model pages, and direct follow-up papers using the model.",
    },
    {
        "Step": 2,
        "Search Area": "Citation-style and Scholar-like discovery",
        "Queries / Methods": '"dascoli2026tribe"; "d\'Ascoli" "TRIBE v2"; "TRIBE v2" "Google Scholar"; author profiles; Semantic Scholar; DBLP',
        "Purpose": "Find papers that cite or reuse TRIBE v2 but may not appear in generic web search.",
        "Inclusion Rule": "Include scholarly records, preprints, and named Google Scholar profile hits when enough metadata is available.",
    },
    {
        "Step": 3,
        "Search Area": "Use as synthetic fMRI / encoding-decoding cascade",
        "Queries / Methods": '"TRIBE v2" "synthetic fMRI"; "TRIBE v2 Data Augmentation"; "encoding model generated synthetic neural data"; "brain simulator decoder training"',
        "Purpose": "Find studies closest to your proposed video -> TRIBE fMRI -> decoder strategy.",
        "Inclusion Rule": "Include direct TRIBE synthetic-fMRI augmentation and adjacent non-TRIBE synthetic-fMRI papers only as context.",
    },
    {
        "Step": 4,
        "Search Area": "Affective neuroscience and emotion",
        "Queries / Methods": '"TRIBE v2" "EmoMind"; "TRIBE v2" emotion; "TRIBE v2" affective fMRI; "TRIBE v2" "Emo-FilM"',
        "Purpose": "Find evidence for or against using TRIBE v2 for emotion-category, valence/arousal, or affective-caption decoding.",
        "Inclusion Rule": "Include TRIBE substitution probes and emotion fMRI papers that directly test synthetic TRIBE fMRI.",
    },
    {
        "Step": 5,
        "Search Area": "Failure, negative, and out-of-domain results",
        "Queries / Methods": '"TRIBE v2" "does not"; "TRIBE" "does not predict"; "TRIBE v2" "limitations"; "subject-agnostic augmentation has limits"',
        "Purpose": "Find studies that contradict over-strong claims: engagement, individualized affect, diagnosis, behavior.",
        "Inclusion Rule": "Include null results, explicit limitations, and studies where TRIBE-derived signals degrade or fail.",
    },
    {
        "Step": 6,
        "Search Area": "Medical / clinical uses",
        "Queries / Methods": '"TRIBE v2" clinical; diagnosis; patient; medical; depression; autism; PTSD; Alzheimer; schizophrenia; stroke',
        "Purpose": "Determine whether TRIBE v2 has been validated in medical diagnosis or patient prediction.",
        "Inclusion Rule": "Separate true studies from blog/social claims; mark clinical status conservatively.",
    },
    {
        "Step": 7,
        "Search Area": "Low-confidence applied papers",
        "Queries / Methods": 'site:ssrn.com "TRIBE v2"; "TRIBE v2 In-Silico Study"; "TRIBE v2 In-Silico fMRI Analysis"; "Journal of Cortexplore" "TRIBE v2"',
        "Purpose": "Capture applied white papers and speculative uses while avoiding treating them as strong validation.",
        "Inclusion Rule": "Include with low-confidence / non-peer-reviewed labels and explicit limitations.",
    },
]


studies = [
    {
        "ID": "S01",
        "Study": "A foundation model of vision, audition, and language for in-silico neuroscience",
        "Authors": "d'Ascoli, Rapin, Benchetrit, Brooks, Begany, Raugel, Banville, King",
        "Year": 2026,
        "Type": "Primary TRIBE v2 paper / arXiv",
        "Direct TRIBE v2 Use": "Introduces TRIBE v2, a tri-modal video/audio/language model predicting fMRI responses.",
        "What They Did": "Unified >1,000 hours of fMRI across 720 subjects; trained a transformer-based multimodal encoder to predict high-resolution cortical and subcortical responses; tested zero-shot generalization and in-silico localizer-style paradigms.",
        "Main Findings": "Reported strong gains over traditional linear encoding models; recovered established visual and neurolinguistic effects in silico; model supports high-resolution average-subject prediction for new stimuli, tasks, and subjects.",
        "Where It Falls Short": "Not a clinical or diagnostic model; public inference is mainly average-subject; fMRI temporal resolution is slow; subjective emotion, symptoms, diagnosis, behavior, and individual differences are not directly validated.",
        "Relevance to Your Project": "Best source for generating synthetic fMRI from CK emotional videos, but must be validated against Kamitani real fMRI and CK labels.",
        "Evidence Direction": "Supports use as an encoding/simulation prior, not as standalone emotion or diagnosis truth.",
        "Medical Context": "Only broad clinical motivation, no patient/diagnosis validation.",
        "Confidence": "High",
        "Peer Review Status": "arXiv / Meta research publication",
        "URL": "https://arxiv.org/abs/2605.04326",
        "Secondary URL": "https://ai.meta.com/research/publications/a-foundation-model-of-vision-audition-and-language-for-in-silico-neuroscience/",
    },
    {
        "ID": "S02",
        "Study": "TRIBE: TRImodal Brain Encoder for whole-brain fMRI response prediction",
        "Authors": "d'Ascoli, Rapin, Benchetrit, Banville, King",
        "Year": 2025,
        "Type": "Predecessor TRIBE / Algonauts 2025 technical report",
        "Direct TRIBE v2 Use": "Predecessor, not TRIBE v2; architecture basis for v2.",
        "What They Did": "Combined text, audio, and video foundation features with a transformer to predict whole-brain fMRI responses to movies in the Algonauts 2025 challenge.",
        "Main Findings": "Won Algonauts 2025; multimodal model outperformed unimodal models, especially in high-level associative cortices.",
        "Where It Falls Short": "Lower-resolution challenge setting; few subjects; not designed for emotion labels, diagnosis, or synthetic-data augmentation.",
        "Relevance to Your Project": "Justifies transformer multimodal encoding for naturalistic video; useful as historical baseline and design rationale.",
        "Evidence Direction": "Supports multimodal encoding over vision-only for emotional videos.",
        "Medical Context": "None.",
        "Confidence": "High",
        "Peer Review Status": "arXiv / OpenReview challenge report",
        "URL": "https://arxiv.org/abs/2507.22229",
        "Secondary URL": "https://github.com/facebookresearch/algonauts-2025",
    },
    {
        "ID": "S03",
        "Study": "Boosting Brain-to-Image Decoding with TRIBE v2 Data Augmentation",
        "Authors": "Benchetrit, Careil, Dahan, Banville, d'Ascoli, King",
        "Year": 2026,
        "Type": "Direct synthetic-fMRI augmentation study",
        "Direct TRIBE v2 Use": "Used TRIBE v2 to generate synthetic fMRI for images and augment fMRI-to-image decoders.",
        "What They Did": "Converted static images into short videos, generated TRIBE synthetic fMRI, mixed synthetic and real fMRI for NSD and BOLD5000 decoder training, evaluated only on held-out real fMRI.",
        "Main Findings": "Reported up to 68% improvement in Top-10 image retrieval in low-data regimes; synthetic-only training could be above chance in some settings; DynaDiff reconstruction metrics improved in a preliminary test.",
        "Where It Falls Short": "Not plug-and-play; too much synthetic data can saturate or hurt; optimal synthetic:real ratio differs by dataset/decoder; public TRIBE output is subject-agnostic; static images are out-of-distribution for a video-trained model.",
        "Relevance to Your Project": "Closest precedent for your video -> TRIBE fMRI -> decoder pipeline. Use TRIBE as augmentation/pretraining, not as sole evidence.",
        "Evidence Direction": "Strongly supports synthetic TRIBE fMRI as a low-data prior, with calibration.",
        "Medical Context": "None; image decoding only.",
        "Confidence": "High",
        "Peer Review Status": "arXiv preprint",
        "URL": "https://arxiv.org/abs/2606.06345",
        "Secondary URL": "https://arxiv.org/html/2606.06345v1",
    },
    {
        "ID": "S04",
        "Study": "EmoMind: Decoding Affective Captions from Human Brain fMRI",
        "Authors": "Mohammed, Gu, Fang",
        "Year": 2026,
        "Type": "Affective fMRI decoding with TRIBE substitution probe",
        "Direct TRIBE v2 Use": "Used TRIBE v2 substitution as a synthetic-brain probe for affective-caption decoding.",
        "What They Did": "Decoded continuous 34-dimensional emotion vectors from real fMRI, generated affective captions, compared against prompted GPT-4, and substituted TRIBE-predicted fMRI for real fMRI in a robustness/geometry test.",
        "Main Findings": "Continuous brain-decoded affect outperformed label-prompted GPT-4 on subject-specificity, structural geometry, and causal-control metrics. The TRIBE substitution preserved some target-conditional affect but sharply weakened relational/structural affect geometry.",
        "Where It Falls Short": "TRIBE synthetic responses did not preserve full person-specific affective structure; real fMRI remained necessary for individualized affective organization.",
        "Relevance to Your Project": "Most directly affective-neuroscience-relevant warning: TRIBE alone may decode coarse emotion but can lose relational/individual affect geometry.",
        "Evidence Direction": "Mixed: supports TRIBE as probe/augmentation, contradicts using TRIBE alone as a substitute for real affective fMRI.",
        "Medical Context": "Mental-health relevance discussed conceptually; no clinical validation.",
        "Confidence": "High",
        "Peer Review Status": "arXiv preprint",
        "URL": "https://arxiv.org/abs/2605.16739",
        "Secondary URL": "https://arxiv.org/html/2605.16739v1",
    },
    {
        "ID": "S05",
        "Study": "Feature Visualization Recovers Known Cortical Selectivity from TRIBE v2",
        "Authors": "Bladon, Bent",
        "Year": 2026,
        "Type": "Interpretability / feature visualization",
        "Direct TRIBE v2 Use": "Optimized images through V-JEPA 2 + TRIBE v2 to maximize predicted ROI activity.",
        "What They Did": "Applied gradient-ascent feature visualization to seven visual ROIs: V1, V2, V3, V4, MT, FFA, and PPA.",
        "Main Findings": "Recovered V1-to-V4 feature hierarchy, MT-like implied-motion streaks, FFA face-like features, and PPA rectilinear/place-like patterns.",
        "Where It Falls Short": "Qualitative interpretability; optimized FFA stimuli were adversarial super-stimuli; no real human validation of generated stimuli; no emotion or clinical decoding.",
        "Relevance to Your Project": "Useful for validating whether TRIBE emotion maps are anatomically plausible, but not sufficient for affective claims.",
        "Evidence Direction": "Supports interpretability and ROI sanity checks.",
        "Medical Context": "None.",
        "Confidence": "High",
        "Peer Review Status": "arXiv preprint",
        "URL": "https://arxiv.org/abs/2605.13904",
        "Secondary URL": "https://arxiv.org/html/2605.13904v1",
    },
    {
        "ID": "S06",
        "Study": "A global predicted-fMRI drive signal from TRIBE does not predict YouTube replay heatmaps",
        "Authors": "Sahu, Pandey",
        "Year": 2026,
        "Type": "Negative/null behavioral prediction study",
        "Direct TRIBE v2 Use": "Ran TRIBE/Algonauts-style predicted fMRI on YouTube videos and reduced cortical responses to engagement curves.",
        "What They Did": "Compared per-second predicted cortical global field power and network/ROI readouts against YouTube most-replayed heatmaps for 48 videos; tested controls and learned readouts.",
        "Main Findings": "No evidence that predicted fMRI global drive predicted re-watch behavior. Position-controlled partial correlation was near zero; null held across six network readouts and permutation tests.",
        "Where It Falls Short": "Sample was restricted to already-popular videos; engagement proxy was behavioral replay rather than emotion; public average-subject model lacks per-subject responses; nucleus accumbens/ventral striatum absent from cortical surface tests.",
        "Relevance to Your Project": "Strong caution against using TRIBE-derived activation as direct evidence of engagement, preference, diagnosis, or behavioral outcome.",
        "Evidence Direction": "Contradicts overclaim that TRIBE predicted fMRI directly predicts behavior.",
        "Medical Context": "None.",
        "Confidence": "High",
        "Peer Review Status": "arXiv preprint",
        "URL": "https://arxiv.org/abs/2607.01400",
        "Secondary URL": "https://arxiv.org/html/2607.01400v2",
    },
    {
        "ID": "S07",
        "Study": "Neurological Plausibility of AI-Generated Music for Commercial Environments: An In-Silico Cortical Investigation Using Wubble and TRIBE v2",
        "Authors": "Sufi",
        "Year": 2026,
        "Type": "Applied in-silico audio/music preprint",
        "Direct TRIBE v2 Use": "Audio-only TRIBE v2 inference on AI-generated instrumental tracks.",
        "What They Did": "Generated five prompt-conditioned music tracks varying arousal, density, and valence; summarized TRIBE predictions in auditory, temporal, temporo-parietal, and inferior frontal HCP parcels.",
        "Main Findings": "Fast bright major-pop condition produced the largest whole-cortex mean and strongest prefrontal composite response; pairwise spatial correlations suggested prompt variation changed predicted cortical states.",
        "Where It Falls Short": "No human listeners, no behavioral outcomes, no real fMRI validation, no subcortical reward/amygdala inference; commercial rather than clinical/affective neuroscience validation.",
        "Relevance to Your Project": "Shows how to structure a reproducible in-silico stimulus-family analysis, but reviewers will require real-fMRI validation for emotional videos.",
        "Evidence Direction": "Supports hypothesis-screening use; weak evidence for actual affective response.",
        "Medical Context": "None.",
        "Confidence": "Medium",
        "Peer Review Status": "arXiv preprint",
        "URL": "https://arxiv.org/abs/2604.04025",
        "Secondary URL": "https://arxiv.org/html/2604.04025v1",
    },
    {
        "ID": "S08",
        "Study": "Toward an In-Silico Neuroscience of Intelligence Analysis: TRIBE v2 Predictions of Stress and Persuasion in Intelligence-Grade Media",
        "Authors": "Begcecanli, Arslan",
        "Year": 2026,
        "Type": "Applied SSRN preprint / non-peer-reviewed",
        "Direct TRIBE v2 Use": "Applied TRIBE v2 to intelligence-grade and politically charged media stimuli.",
        "What They Did": "Used TRIBE v2 predicted cortical responses to compare stress/persuasion-like responses to media categories.",
        "Main Findings": "Preliminary claim that intelligence-relevant information types do not produce identical predicted neural profiles.",
        "Where It Falls Short": "Non-peer-reviewed; details accessible mainly through SSRN/search snippets; no real fMRI validation; likely high risk of reverse-inference from model-predicted cortical maps.",
        "Relevance to Your Project": "Example of using TRIBE for affect/stress-like media analysis, but it is a cautionary model of how not to overclaim without human validation.",
        "Evidence Direction": "Weak support for in-silico media screening; weak scientific evidence.",
        "Medical Context": "Stress/persuasion only; no clinical diagnosis.",
        "Confidence": "Low",
        "Peer Review Status": "SSRN preprint; listed as non-peer-reviewed on author profile",
        "URL": "https://papers.ssrn.com/sol3/papers.cfm?abstract_id=6726542",
        "Secondary URL": "https://avesis.pa.edu.tr/alpcenk.arslan/yayinlar",
    },
    {
        "ID": "S09",
        "Study": "Predicted Cortical Signatures of Surveillance, Comfort, and Exit Rights in a Speculative Network State Scenario: A TRIBE v2 In-Silico Study",
        "Authors": "Begcecanli",
        "Year": 2026,
        "Type": "Applied SSRN preprint / speculative in-silico study",
        "Direct TRIBE v2 Use": "Used TRIBE v2 to estimate cortical responses to speculative governance/network-state stimuli.",
        "What They Did": "Compared model-predicted cortical signatures for surveillance, comfort, and exit-rights scenarios.",
        "Main Findings": "Claimed TRIBE v2 can provide predicted cortical signatures for abstract sociopolitical stimuli.",
        "Where It Falls Short": "Speculative; no human validation; no medical context; abstract constructs may exceed what a stimulus-driven fMRI encoder can support.",
        "Relevance to Your Project": "Cautionary boundary case: TRIBE can generate maps for any stimulus, but interpretation may be much weaker than prediction.",
        "Evidence Direction": "Mostly cautionary / low-confidence applied use.",
        "Medical Context": "None.",
        "Confidence": "Low",
        "Peer Review Status": "SSRN preprint",
        "URL": "https://papers.ssrn.com/sol3/papers.cfm?abstract_id=6795258",
        "Secondary URL": "https://papers.ssrn.com/sol3/Delivery.cfm/6795258.pdf?abstractid=6795258&mirid=1",
    },
    {
        "ID": "S10",
        "Study": "Decoding Consumer Neural Response to kalories Packaging: A TRIBE v2 In-Silico fMRI Analysis",
        "Authors": "Gandla, Jallepalli",
        "Year": 2026,
        "Type": "White paper / Journal of Cortexplore",
        "Direct TRIBE v2 Use": "Used TRIBE v2 on a product packaging video.",
        "What They Did": "Predicted cortical responses to a dark-chocolate packaging video; identified temporal peaks in mean cortical activation.",
        "Main Findings": "Reported highest mean cortical activation at 21 s and a secondary cluster at 56-58 s; framed these as cognitive-emotional intensity peaks for marketing optimization.",
        "Where It Falls Short": "White paper, product-specific, no real fMRI, no behavior, no held-out validation; wording treats synthetic model output as empirical evidence, which is scientifically risky.",
        "Relevance to Your Project": "Useful warning: for your paper, do not call TRIBE-only results 'empirical human fMRI'; validate against Kamitani real fMRI.",
        "Evidence Direction": "Cautionary; method demonstration but weak scientific support.",
        "Medical Context": "None; neuromarketing.",
        "Confidence": "Low",
        "Peer Review Status": "Journal white paper; review rigor unclear",
        "URL": "https://cortexplore.org/index.php/jce/article/view/12",
        "Secondary URL": "https://www.researchgate.net/publication/404701316_Decoding_Consumer_Neural_Response_to_kalories_Packaging_A_TRIBE_v2_In-Silico_fMRI_Analysis",
    },
]


medical_findings = [
    {
        "Finding": "No direct TRIBE v2 clinical validation study found",
        "Evidence": "Searches for TRIBE v2 with clinical, diagnosis, patient, depression, autism, PTSD, Alzheimer, schizophrenia, and stroke returned official/blog claims and social commentary, not patient-level validation papers.",
        "Implication": "Do not claim diagnosis from TRIBE v2 synthetic fMRI alone. A valid clinical study needs real patient fMRI/EEG or symptom data.",
        "Source": "Search audit across Google/Scholar-like queries on 2026-07-21",
        "URL": "https://ai.meta.com/blog/tribe-v2-brain-predictive-foundation-model/",
    },
    {
        "Finding": "Official Meta materials mention clinical researchers and neurological disorders",
        "Evidence": "Meta states TRIBE v2 may help researchers test theories without human subjects and accelerate research toward neurological-disorder treatments.",
        "Implication": "This is motivation, not evidence of diagnostic performance.",
        "Source": "Meta TRIBE v2 blog",
        "URL": "https://ai.meta.com/blog/tribe-v2-brain-predictive-foundation-model/",
    },
    {
        "Finding": "Average-subject predictions limit diagnosis claims",
        "Evidence": "Public TRIBE v2 inference predicts average-subject cortical activity; a negative replay-heatmap paper notes released facebook/tribev2 cannot supply subject-specific responses for some tests.",
        "Implication": "Clinical decoding should use real patient-minus-normative residuals, subject adaptation, or patient-specific data, not synthetic average responses alone.",
        "Source": "TRIBE v2 GitHub and Sahu/Pandey null study",
        "URL": "https://github.com/facebookresearch/tribev2",
    },
    {
        "Finding": "Emotion/mental-health translation remains plausible but unvalidated",
        "Evidence": "EmoMind frames individualized affective fMRI as relevant to mental health, but its TRIBE substitution probe indicates loss of person-specific affective geometry.",
        "Implication": "For depression/PTSD/autism work, test whether real-subject residuals from a normative TRIBE model predict symptoms.",
        "Source": "EmoMind",
        "URL": "https://arxiv.org/abs/2605.16739",
    },
]


project_takeaways = [
    {
        "Priority": 1,
        "Recommendation": "Use TRIBE v2 as a normative encoding prior, not as ground truth.",
        "Rationale": "The strongest direct augmentation paper shows gains only when synthetic and real fMRI are mixed and calibrated; EmoMind and YouTube replay results show TRIBE signals can lose affective geometry or fail behavior prediction.",
        "Concrete Next Step": "Run TRIBE on the same CK videos, then compare TRIBE RDMs/maps/decoders against Kamitani real fMRI before training emotion decoders.",
    },
    {
        "Priority": 2,
        "Recommendation": "Make the main scientific question validation-focused.",
        "Rationale": "Kamitani already studied visually evoked emotion. Your novelty is whether TRIBE reproduces or fails to reproduce real affective fMRI organization.",
        "Concrete Next Step": "Paper question: Can TRIBE v2 simulate high-dimensional affective brain representations during emotional video viewing?",
    },
    {
        "Priority": 3,
        "Recommendation": "Separate emotion-category decoding from diagnosis.",
        "Rationale": "TRIBE synthetic average responses may classify video emotion, but diagnosis requires subject-specific patient deviations.",
        "Concrete Next Step": "For diagnosis, model real fMRI - TRIBE normative fMRI residuals and predict symptoms/diagnosis only if patient data exist.",
    },
    {
        "Priority": 4,
        "Recommendation": "Use strong baselines.",
        "Rationale": "A TRIBE -> label decoder may just be a video classifier through a brain-shaped bottleneck.",
        "Concrete Next Step": "Compare direct video features, VLM captions, CK labels, real fMRI, TRIBE fMRI, and TRIBE+real augmentation.",
    },
    {
        "Priority": 5,
        "Recommendation": "Emphasize failure modes as part of the paper.",
        "Rationale": "Negative evidence is valuable: TRIBE may capture visual/semantic structure better than subjective affect, engagement, or clinical individuality.",
        "Concrete Next Step": "Report where TRIBE matches real fMRI and where it diverges: visual cortex vs salience/default/affective regions, category vs valence/arousal, group vs subject-specific structure.",
    },
]


search_log = [
    ["Exact phrases", '"TRIBE v2" fMRI; "TRIBEv2" fMRI; "TRIBE v2" "synthetic fMRI"; "TRIBE v2" "arXiv"'],
    ["Citation strings", '"dascoli2026tribe"; "d\'Ascoli" "TRIBE v2"; "facebook/tribev2" "arxiv"'],
    ["Affective terms", '"TRIBE v2" EmoMind; "TRIBE v2" emotion; "TRIBE v2" affective; "TRIBE v2" Emo-FilM'],
    ["Failure terms", '"TRIBE v2" "does not"; "TRIBE" "does not predict"; "subject-agnostic augmentation has limits"'],
    ["Medical terms", '"TRIBE v2" clinical; diagnosis; patient; depression; autism; PTSD; Alzheimer; schizophrenia; stroke'],
    ["Applied/grey literature", 'site:ssrn.com "TRIBE v2"; "TRIBE v2 In-Silico Study"; "Journal of Cortexplore" "TRIBE v2"'],
]


def write_sheet(ws, rows):
    if not rows:
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_table(ws)


def style_table(ws):
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    border = Border(bottom=thin)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = border
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for col_idx, col in enumerate(ws.columns, start=1):
        header = ws.cell(1, col_idx).value or ""
        max_len = len(str(header))
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 80))
        width = min(max(max_len + 2, 12), 55)
        if header in {"What They Did", "Main Findings", "Where It Falls Short", "Relevance to Your Project", "Rationale", "Concrete Next Step", "Evidence"}:
            width = 45
        if header in {"URL", "Secondary URL"}:
            width = 48
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 72
    ws.row_dimensions[1].height = 34


def add_summary_sheet(wb):
    ws = wb.create_sheet("Executive Summary", 0)
    rows = [
        ["TRIBE v2 Literature Audit", "Updated 2026-07-21"],
        ["Scope", "Direct TRIBE v2 studies, predecessor TRIBE, failure/negative evidence, medical/clinical search results, and project implications for CK emotional videos + Kamitani fMRI."],
        ["Bottom Line", "TRIBE v2 is promising as a synthetic fMRI generator and normative encoding prior, but current literature does not validate it as a standalone emotion, behavior, or diagnosis model."],
        ["Direct studies found", len(studies)],
        ["High-confidence direct studies", sum(1 for s in studies if s["Confidence"] == "High")],
        ["Medical TRIBE v2 validation studies found", 0],
        ["Most supportive evidence", "Benchetrit et al. 2026: TRIBE synthetic fMRI improved image decoding in low-data regimes, but required calibration."],
        ["Strongest caution", "EmoMind: TRIBE substitution weakened affective relational geometry; Sahu/Pandey: TRIBE predicted fMRI did not predict YouTube rewatch heatmaps."],
        ["Best strategy for your project", "Validate TRIBE against real Kamitani fMRI first; then use TRIBE for augmentation, normative residuals, ROI/RSA comparisons, and ablation-controlled decoding."],
    ]
    for r in rows:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["B1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E79")
    ws["B1"].fill = PatternFill("solid", fgColor="1F4E79")
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    for i in range(2, ws.max_row + 1):
        ws.row_dimensions[i].height = 45


def add_search_log_sheet(wb):
    ws = wb.create_sheet("Search Log")
    ws.append(["Category", "Queries"])
    for row in search_log:
        ws.append(row)
    style_table(ws)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    add_summary_sheet(wb)
    ws = wb.create_sheet("Search Plan")
    write_sheet(ws, search_plan)
    add_search_log_sheet(wb)
    ws = wb.create_sheet("Study Matrix")
    write_sheet(ws, studies)
    ws = wb.create_sheet("Negative Evidence")
    negative = [s for s in studies if "Contradicts" in s["Evidence Direction"] or "warning" in s["Relevance to Your Project"].lower() or s["ID"] in {"S04", "S06", "S10"}]
    write_sheet(ws, negative)
    ws = wb.create_sheet("Medical Clinical")
    write_sheet(ws, medical_findings)
    ws = wb.create_sheet("Project Takeaways")
    write_sheet(ws, project_takeaways)
    wb.save(OUT_XLSX)

    # Verification pass: reopen and check sheet names / row counts.
    check = load_workbook(OUT_XLSX, read_only=True, data_only=True)
    expected = {
        "Executive Summary": 9,
        "Search Plan": len(search_plan) + 1,
        "Search Log": len(search_log) + 1,
        "Study Matrix": len(studies) + 1,
        "Negative Evidence": len(negative) + 1,
        "Medical Clinical": len(medical_findings) + 1,
        "Project Takeaways": len(project_takeaways) + 1,
    }
    actual = {name: check[name].max_row for name in expected}
    if actual != expected:
        raise RuntimeError(f"Verification failed: {actual} != {expected}")
    print(str(OUT_XLSX))
    print(actual)


if __name__ == "__main__":
    main()
